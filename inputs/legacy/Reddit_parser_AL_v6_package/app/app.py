"""
Reddit Research Tool — локальний веб-інтерфейс для пошуку, підрахунку
і збору постів+коментарів з Reddit.

Запуск:  py app.py
Потім відкрити в браузері: http://localhost:5000
"""

import re
import time
import json
import html
import threading
import shutil
import traceback
from pathlib import Path
from datetime import datetime, timezone

import requests
from flask import Flask, request, jsonify, send_file, abort
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
SETTINGS_FILE = BASE_DIR / "saved_settings.json"

DEFAULT_SETTINGS = {
    "base_name": "reddit_project",
    "mode": "count",
    "keywords": 'ABM\n"ABM strategy"',
    "subreddits": "marketing",
    "sort": "new",
    "time_filter": "all",
    "months_back": 0,
    "limit": 1000,
    "max_depth": 6,
    "min_score": -10,
    "delay": 5,
    "posts_per_file": 120,
}


def load_settings() -> dict:
    """Завантажує налаштування останнього запуску. Cookie тут ніколи не зберігається."""
    if SETTINGS_FILE.exists():
        try:
            saved = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
            merged = dict(DEFAULT_SETTINGS)
            merged.update(saved)
            return merged
        except Exception:
            pass
    return dict(DEFAULT_SETTINGS)


def save_settings(params: dict):
    """Зберігає налаштування для наступного запуску. Cookie свідомо не включається."""
    to_save = {k: params[k] for k in DEFAULT_SETTINGS.keys() if k in params}
    to_save["keywords"] = "\n".join(params.get("keywords", []))
    to_save["subreddits"] = "\n".join(params.get("subreddits", []))
    try:
        SETTINGS_FILE.write_text(json.dumps(to_save, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as e:
        log(f"⚠ Не вдалось зберегти налаштування: {e}")

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "en-US,en;q=0.9",
}

# ══════════════════════════════════════════════════════════
# ГЛОБАЛЬНИЙ СТАН ЗАДАЧІ (один прогін одночасно)
# ══════════════════════════════════════════════════════════

job_lock = threading.Lock()
job = {
    "status": "idle",       # idle / running / paused / stopping / done / error
    "log": [],
    "current": 0,
    "total": 0,
    "comments": 0,           # опрацьовано коментарів (живий рахунок)
    "result_files": [],     # список файлів для завантаження
    "mode": None,
}
pause_event = threading.Event()
stop_event = threading.Event()
session_cookie = {"reddit_session": ""}  # мутабельний — можна оновити без перезапуску задачі
consecutive_403 = {"count": 0}


def log(msg: str):
    ts = datetime.now().strftime("%H:%M:%S")
    with job_lock:
        job["log"].append(f"[{ts}] {msg}")
        if len(job["log"]) > 2000:
            job["log"] = job["log"][-2000:]


def set_progress(current, total):
    with job_lock:
        job["current"] = current
        job["total"] = total


def add_comments(n: int):
    with job_lock:
        job["comments"] += n


def wait_if_paused():
    """Блокує виконання, поки стоїть пауза. Повертає True, якщо треба зупинитись повністю."""
    if pause_event.is_set():
        log("⏸ Пауза...")
        with job_lock:
            job["status"] = "paused"
    while pause_event.is_set():
        if stop_event.is_set():
            return True
        time.sleep(0.4)
    if stop_event.is_set():
        return True
    with job_lock:
        if job["status"] == "paused":
            job["status"] = "running"
    return False


# ══════════════════════════════════════════════════════════
# ПОШУК (спільна логіка для підрахунку і повного збору)
# ══════════════════════════════════════════════════════════

def build_search_url(search_term, subreddit, sort, time_filter, page_size, after=None):
    if subreddit:
        base = f"https://www.reddit.com/r/{subreddit}/search.json"
        params = f"q={requests.utils.quote(search_term)}&restrict_sr=1&sort={sort}&t={time_filter}&limit={page_size}"
    else:
        base = "https://www.reddit.com/search.json"
        params = f"q={requests.utils.quote(search_term)}&sort={sort}&t={time_filter}&limit={page_size}"
    url = f"{base}?{params}"
    if after:
        url += f"&after={after}"
    return url


def search_paginated(query, subreddit, sort, time_filter, limit, months_back, collect_permalinks=True):
    """
    Проходить пагінацію пошуку. Повертає (permalinks_list, count, error_or_None).
    Якщо collect_permalinks=False — не зберігає посилання, тільки рахує (швидше для матриці).
    """
    # Точний/широкий пошук тепер контролює сам користувач: якщо він написав
    # ключове слово в лапках ("ABM strategy") — це піде як точна фраза.
    # Без лапок (ABM strategy) Reddit шукає нечітко/широко. Скрипт нічого
    # не додає й не забирає — бере запит рядок як є.
    search_term = query.strip()
    cutoff_ts = time.time() - (months_back * 30 * 24 * 60 * 60) if months_back else None

    urls = []
    count = 0
    after = None

    while count < limit:
        if wait_if_paused():
            return urls, count, "stopped"

        page_size = min(100, limit - count)
        url = build_search_url(search_term, subreddit, sort, time_filter, page_size, after)

        try:
            resp = requests.get(url, headers=HEADERS, cookies=session_cookie, timeout=15)
        except requests.exceptions.RequestException as e:
            return urls, count, f"Помилка з'єднання: {e}"

        if resp.status_code == 403:
            return urls, count, "403 — cookie не діє або застаріла"
        if resp.status_code == 429:
            log("⚠ Rate limit на пошуку, чекаємо 60 сек...")
            for _ in range(4):
                time.sleep(15)
                log("⏳ ...ще чекаю (це нормально, не зависання)")
            continue
        if resp.status_code != 200:
            return urls, count, f"HTTP {resp.status_code}"

        try:
            data = resp.json()
        except Exception as e:
            return urls, count, f"Помилка JSON: {e}"

        children = data.get("data", {}).get("children", [])
        if not children:
            break

        hit_older = False
        for item in children:
            d = item.get("data", {})
            created_utc = d.get("created_utc", 0)
            if cutoff_ts and created_utc < cutoff_ts:
                hit_older = True
                break
            count += 1
            if collect_permalinks:
                permalink = d.get("permalink")
                if permalink:
                    urls.append("https://www.reddit.com" + permalink)
            if count >= limit:
                break

        if hit_older or count >= limit:
            break

        after = data.get("data", {}).get("after")
        if not after:
            break

        time.sleep(1)

    return urls, count, None


# ══════════════════════════════════════════════════════════
# РЕЖИМ 1 — ТІЛЬКИ ПІДРАХУНОК (МАТРИЦЯ)
# ══════════════════════════════════════════════════════════

def run_count_job(params):
    subreddits = params["subreddits"]
    keywords = params["keywords"]
    sort = params["sort"]
    time_filter = params["time_filter"]
    limit = params["limit"]
    months_back = params["months_back"]
    base_name = params["base_name"]

    total = len(subreddits) * len(keywords)
    set_progress(0, total)
    results = []
    i = 0

    for sub in subreddits:
        for kw in keywords:
            if stop_event.is_set():
                log("⏹ Зупинено користувачем.")
                break
            if wait_if_paused():
                log("⏹ Зупинено користувачем.")
                break

            i += 1
            log(f"[{i}/{total}] r/{sub} — \"{kw}\"")

            _, count, error = search_paginated(
                kw, sub, sort, time_filter, limit, months_back,
                collect_permalinks=False,
            )

            if error == "stopped":
                break
            elif error:
                log(f"    ✗ {error}")
                results.append((sub, kw, "ПОМИЛКА"))
            else:
                capped = " (можливо більше)" if count >= limit else ""
                log(f"    → знайдено: {count}{capped}")
                results.append((sub, kw, count))

            set_progress(i, total)
            time.sleep(2)
        if stop_event.is_set():
            break

    # ── Зберігаємо у справжній Excel ──
    out_path = OUTPUT_DIR / f"{base_name}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Матриця"

    headers = ["Сабредіт", "Ключове слово", "Кількість постів"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Arial", color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    for row in results:
        ws.append(list(row))

    for col_idx, header in enumerate(headers, start=1):
        values = [str(header)] + [str(r[col_idx - 1]) for r in results]
        max_len = max(len(v) for v in values)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    ws.freeze_panes = "A2"
    if results:
        ws.auto_filter.ref = f"A1:C{len(results) + 1}"

    wb.save(out_path)
    log(f"✅ Готово. Таблицю збережено: {out_path.name}")

    with job_lock:
        job["result_files"] = [out_path.name]
        job["status"] = "stopped" if stop_event.is_set() else "done"


# ══════════════════════════════════════════════════════════
# РЕЖИМ 2 — ПОВНИЙ ЗБІР (ТЕКСТ + КОМЕНТАРІ)
# ══════════════════════════════════════════════════════════

def extract_post_id(url: str) -> str:
    m = re.search(r'/comments/([a-z0-9]+)', url)
    return m.group(1) if m else ""


def load_done_ids(out_dir: Path) -> set:
    done = set()
    for p in out_dir.glob("posts_part_*.md"):
        try:
            text = p.read_text(encoding="utf-8")
        except Exception:
            continue
        for m in re.finditer(r'<!-- id:([a-z0-9]+) -->', text):
            done.add(m.group(1))
    return done


def fetch_post_with_comments(url, comment_limit, max_retries, retries=0):
    """Повертає (data, error_reason). data=None якщо не вдалось; error_reason='403'/'404'/... для діагностики."""
    clean_url = re.sub(r'\?.*$', '', url.rstrip("/"))
    json_url = clean_url + f".json?limit={comment_limit}"

    try:
        response = requests.get(json_url, headers=HEADERS, cookies=session_cookie, timeout=15)
    except requests.exceptions.RequestException as e:
        log(f"  ✗ Помилка з'єднання: {e}")
        return None, "connection"

    if response.status_code == 200:
        try:
            data = response.json()
            if isinstance(data, list) and len(data) >= 2:
                return data, None
            log("  ✗ Неочікувана структура відповіді")
            return None, "structure"
        except Exception as e:
            log(f"  ✗ Помилка JSON: {e}")
            return None, "json"
    elif response.status_code == 429:
        if retries >= max_retries:
            log("  ✗ Rate limit не знявся після кількох спроб")
            return None, "429"
        log(f"  ⚠ Rate limit! Чекаємо 60 сек... (спроба {retries + 1})")
        for _ in range(4):
            time.sleep(15)
            log("  ⏳ ...ще чекаю (це нормально, не зависання)")
        return fetch_post_with_comments(url, comment_limit, max_retries, retries + 1)
    elif response.status_code == 404:
        log("  ✗ Пост не знайдено або видалено")
        return None, "404"
    elif response.status_code == 403:
        log("  ✗ 403 — доступ заборонено (перевір cookie!)")
        return None, "403"
    else:
        log(f"  ✗ HTTP {response.status_code}")
        return None, "other"


def parse_comments(comments_data, max_depth, min_score, depth=0):
    if depth > max_depth:
        return "", 0
    result = ""
    count = 0
    indent = "  " * depth
    for item in comments_data:
        if not isinstance(item, dict) or item.get("kind") != "t1":
            continue
        d = item.get("data", {})
        body = d.get("body", "").strip()
        score = d.get("score", 0)
        author = d.get("author", "невідомо")
        if body in ("[deleted]", "[removed]", ""):
            continue
        if isinstance(score, int) and score < min_score:
            continue
        result += f"{indent}**u/{author}** (👍 {score})\n{indent}{body}\n\n"
        count += 1
        replies = d.get("replies", "")
        if isinstance(replies, dict):
            children = replies.get("data", {}).get("children", [])
            child_text, child_count = parse_comments(children, max_depth, min_score, depth + 1)
            result += child_text
            count += child_count
    return result, count


def post_to_markdown(raw_data, post_index, search_terms, max_depth, min_score):
    try:
        post_data = raw_data[0]["data"]["children"][0]["data"]
    except (KeyError, IndexError) as e:
        log(f"  ✗ Помилка читання поста: {e}")
        return "", 0

    post_id = post_data.get("id", "")
    title = post_data.get("title", "Без назви")
    author = post_data.get("author", "невідомо")
    subreddit = post_data.get("subreddit_name_prefixed", "r/?")
    body = post_data.get("selftext", "").strip()
    score = post_data.get("score", 0)
    num_comments = post_data.get("num_comments", 0)
    url = post_data.get("url", "")
    permalink = "https://reddit.com" + post_data.get("permalink", "")
    created = datetime.fromtimestamp(post_data.get("created_utc", 0), timezone.utc)
    date_str = created.strftime("%Y-%m-%d %H:%M UTC")

    if body in ("[deleted]", "[removed]"):
        body = "_Текст поста видалено_"
    elif not body:
        body = f"_Пост-посилання: {url}_"

    try:
        comments_raw = raw_data[1]["data"]["children"]
        comments_md, comment_count = parse_comments(comments_raw, max_depth, min_score)
    except (KeyError, IndexError):
        comments_md, comment_count = "", 0

    if not comments_md.strip():
        comments_md = "_Коментарів немає або всі видалено_"

    markdown = f"""
{'='*60}
<!-- id:{post_id} -->
# ПОСТ {post_index}: {title}

- **Знайдено за запитом:** {search_terms}
- **Сабредіт:** {subreddit}
- **Автор:** u/{author}
- **Дата:** {date_str}
- **Upvotes:** {score:,}
- **Коментарів:** {num_comments:,}
- **Посилання:** {permalink}

## Текст поста

{body}

## Коментарі

{comments_md}
"""
    return markdown, comment_count


def write_file_header(filepath, part_number, total_parts, post_from, post_to,
                       total_posts, keywords, subreddits):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"""# REDDIT ПОСТИ — ЧАСТИНА {part_number} з {total_parts}

**Дата експорту:** {datetime.now().strftime("%Y-%m-%d %H:%M")}
**Ключові слова:** {', '.join(keywords)}
**Сабредіти:** {', '.join(subreddits) if subreddits else 'усі'}
**Пости у цьому файлі:** {post_from}–{post_to} з {total_posts}
**Джерело:** Reddit Search JSON API

---

""")


def run_scrape_job(params):
    subreddits = params["subreddits"]
    keywords = params["keywords"]
    sort = params["sort"]
    time_filter = params["time_filter"]
    limit = params["limit"]
    months_back = params["months_back"]
    base_name = params["base_name"]
    max_depth = params["max_depth"]
    min_score = params["min_score"]
    comment_limit = params["comment_limit"]
    delay = params["delay"]
    max_retries = params["max_retries"]
    posts_per_file = params["posts_per_file"]

    out_dir = OUTPUT_DIR / base_name
    out_dir.mkdir(exist_ok=True)

    log("═" * 40)
    log("КРОК 1: Пошук постів за ключовими словами")
    log("═" * 40)

    seen = set()
    ordered_urls = []
    url_keywords = {}  # url -> список ключових слів, якими саме його знайдено
    targets = subreddits if subreddits else [None]

    for query in keywords:
        for sub in targets:
            if stop_event.is_set() or wait_if_paused():
                log("⏹ Зупинено користувачем під час пошуку.")
                with job_lock:
                    job["status"] = "stopped"
                return

            label = f"r/{sub}" if sub else "усі сабредіти"
            log(f"🔎 Пошук: \"{query}\" у {label}")
            found, count, error = search_paginated(
                query, sub, sort, time_filter, limit, months_back,
                collect_permalinks=True,
            )
            if error == "stopped":
                with job_lock:
                    job["status"] = "stopped"
                return
            if error:
                log(f"   ✗ {error}")
                continue
            log(f"   → знайдено {len(found)} посилань")
            for u in found:
                if u not in seen:
                    seen.add(u)
                    ordered_urls.append(u)
                kw_list = url_keywords.setdefault(u, [])
                if query not in kw_list:
                    kw_list.append(query)
            time.sleep(2)

    if not ordered_urls:
        log("❌ Пошук не дав жодного результату.")
        with job_lock:
            job["status"] = "error"
        return

    log(f"✅ Всього унікальних постів знайдено: {len(ordered_urls)}")
    urls_file = out_dir / "search_found_urls.txt"
    urls_file.write_text("\n".join(ordered_urls), encoding="utf-8")

    log("═" * 40)
    log("КРОК 2: Збір повного тексту + коментарів")
    log("═" * 40)

    total = len(ordered_urls)
    total_parts = (total + posts_per_file - 1) // posts_per_file
    done_ids = load_done_ids(out_dir)
    if done_ids:
        log(f"↻ Знайдено вже зібраних постів: {len(done_ids)} — пропускаю їх")

    set_progress(0, total)
    success = 0
    skipped = 0
    failed = []
    current_part = 0
    current_file = None
    file_logged = False
    post_from = post_to = 0

    for i, url in enumerate(ordered_urls, start=1):
        if stop_event.is_set() or wait_if_paused():
            log("⏹ Зупинено користувачем.")
            with job_lock:
                job["status"] = "stopped"
            set_progress(i - 1, total)
            _finalize_scrape(out_dir, base_name)
            return

        # Номер частини визначаємо за позицією посту ще ДО перевірки на пропуск,
        # інакше при повторному запуску (коли перші пости частини вже зібрані)
        # current_file лишився б порожнім.
        part_no = (i - 1) // posts_per_file + 1
        if part_no != current_part:
            current_part = part_no
            post_from = (part_no - 1) * posts_per_file + 1
            post_to = min(part_no * posts_per_file, total)
            current_file = out_dir / f"posts_part_{part_no:03d}_of_{total_parts:03d}.md"
            file_logged = False

        pid = extract_post_id(url)
        if pid and pid in done_ids:
            skipped += 1
            set_progress(i, total)
            continue

        if not current_file.exists():
            write_file_header(current_file, current_part, total_parts,
                               post_from, post_to, total, keywords, subreddits)
        if not file_logged:
            log(f"📄 Файл: {current_file.name} (пости {post_from}–{post_to})")
            file_logged = True

        log(f"[{i}/{total}] {url}")
        raw_data, error_reason = fetch_post_with_comments(url, comment_limit, max_retries)

        if error_reason == "403":
            consecutive_403["count"] += 1
            if consecutive_403["count"] >= 3:
                log("⚠⚠⚠ Три 403-помилки підряд — cookie, ймовірно, застаріла.")
                log("⏸ Автоматична пауза. Встав свіжу cookie у поле нижче і натисни 'Оновити й продовжити'.")
                pause_event.set()
                with job_lock:
                    job["status"] = "paused"
                consecutive_403["count"] = 0
        else:
            consecutive_403["count"] = 0

        if raw_data:
            md_block, comment_count = post_to_markdown(
                raw_data, i,
                ", ".join(url_keywords.get(url, [])),
                max_depth, min_score,
            )
            if md_block:
                with open(current_file, "a", encoding="utf-8") as f:
                    f.write(md_block)
                if pid:
                    done_ids.add(pid)
                success += 1
                add_comments(comment_count)
                log(f"  ✅ Збережено ({comment_count} коментарів)")
            else:
                failed.append(url)
        else:
            failed.append(url)

        set_progress(i, total)
        if i < total:
            time.sleep(delay)

    log(f"✅ Успішно: {success} | Пропущено: {skipped} | Не вдалось: {len(failed)}")

    if failed:
        failed_file = out_dir / "failed_urls.txt"
        failed_file.write_text("\n".join(failed), encoding="utf-8")
        log(f"❌ Список невдалих: {failed_file.name}")

    _finalize_scrape(out_dir, base_name)
    with job_lock:
        job["status"] = "stopped" if stop_event.is_set() else "done"


def _finalize_scrape(out_dir: Path, base_name: str):
    """Пакує папку з результатами в ZIP для зручного завантаження одним файлом."""
    zip_base = OUTPUT_DIR / base_name
    try:
        zip_path = shutil.make_archive(str(zip_base), "zip", root_dir=out_dir)
        with job_lock:
            job["result_files"] = [Path(zip_path).name]
        log(f"📦 Архів готовий: {Path(zip_path).name}")
    except Exception as e:
        log(f"✗ Не вдалось створити ZIP: {e}")


# ══════════════════════════════════════════════════════════
# ЗАПУСК ФОНОВОЇ ЗАДАЧІ
# ══════════════════════════════════════════════════════════

def run_job_thread(params):
    try:
        session_cookie["reddit_session"] = params["cookie"]
        consecutive_403["count"] = 0

        with job_lock:
            job["status"] = "running"
            job["log"] = []
            job["current"] = 0
            job["total"] = 0
            job["comments"] = 0
            job["result_files"] = []
            job["mode"] = params["mode"]

        if params["mode"] == "count":
            run_count_job(params)
        else:
            run_scrape_job(params)

    except Exception:
        log("✗ Помилка виконання:")
        log(traceback.format_exc())
        with job_lock:
            job["status"] = "error"


# ══════════════════════════════════════════════════════════
# FLASK ROUTES
# ══════════════════════════════════════════════════════════

INDEX_HTML = """
<!DOCTYPE html>
<html lang="uk">
<head>
<meta charset="UTF-8">
<title>Reddit Research Tool</title>
<style>
  :root {
    --bg: #101828;
    --card: #1A2438;
    --border: #2E3B57;
    --accent: #B8F500;
    --text: #FFFFFF;
    --text-secondary: #7A8499;
  }
  body { font-family: 'Calibri', 'Segoe UI', sans-serif; max-width: 760px; margin: 30px auto;
    padding: 0 20px; color: var(--text); background: var(--bg); }
  h1 { font-size: 20px; font-weight: 500; }
  fieldset { border: none; background: var(--card); border-radius: 10px; margin-bottom: 16px; padding: 16px 18px; }
  legend { font-weight: 500; color: var(--text); font-size: 14px; padding: 0 0 6px; }
  label { display: block; margin-top: 10px; font-size: 13px; font-weight: 500; color: var(--text); }
  input[type=text], input[type=number], input[type=password], textarea, select {
    width: 100%; padding: 7px; margin-top: 4px; box-sizing: border-box;
    border: 1px solid var(--border); border-radius: 6px; font-size: 13px;
    background: var(--bg); color: var(--text); font-family: 'Calibri', sans-serif;
  }
  textarea { height: 70px; font-family: 'Consolas', 'Calibri', monospace; }
  .row { display: flex; gap: 14px; }
  .row > div { flex: 1; }
  button { background: var(--accent); color: var(--bg); border: none; padding: 12px 22px;
    border-radius: 8px; font-size: 15px; font-weight: 500; font-family: 'Calibri', sans-serif;
    cursor: pointer; margin-top: 16px; }
  button:hover { opacity: 0.9; }
  .hint { color: var(--text-secondary); font-size: 12px; margin-top: 3px; }
  .testbtn { background: var(--border); color: var(--accent); padding: 4px 10px; font-size: 12px;
    margin-left: 8px; margin-top: 4px; vertical-align: middle; }
  .mode-choice label { display: inline-block; margin-right: 20px; font-weight: normal; }
</style>
</head>
<body>
<h1>Reddit Research Tool</h1>
<form id="jobForm">

  <fieldset>
    <legend>Проєкт</legend>
    <label>Назва файлу / проєкту
      <input type="text" name="base_name" value="@@BASE_NAME@@" required>
    </label>
    <div class="hint">Використовується для назви Excel-файлу і папки з постами.</div>

    <label class="mode-choice" style="margin-top:14px;">
      <input type="radio" name="mode" value="count" @@MODE_COUNT_CHECKED@@ onchange="toggleMode()"> Тільки підрахунок (таблиця Excel)
      <input type="radio" name="mode" value="scrape" @@MODE_SCRAPE_CHECKED@@ onchange="toggleMode()"> Повний збір (текст + коментарі)
    </label>
  </fieldset>

  <fieldset>
    <legend>Що шукати</legend>
    <label>Ключові слова (кожне з нового рядка)
      <textarea name="keywords" required>@@KEYWORDS@@</textarea>
    </label>
    <div class="hint">Без лапок — широкий пошук. У лапках ("фраза з кількох слів") — точний збіг фрази.</div>
    <label>Сабредіти (кожен з нового рядка, без "r/". Порожньо = увесь Reddit)
      <textarea name="subreddits">@@SUBREDDITS@@</textarea>
    </label>
  </fieldset>

  <fieldset>
    <legend>Параметри пошуку</legend>
    <div class="row">
      <div>
        <label>Сортування
          <select name="sort">
            @@SORT_OPTIONS@@
          </select>
        </label>
      </div>
      <div>
        <label>Період
          <select name="time_filter">
            @@TIME_OPTIONS@@
          </select>
        </label>
      </div>
    </div>
    <div class="row">
      <div>
        <label>Останні N місяців (0 = без обмеження)
          <input type="number" name="months_back" value="@@MONTHS_BACK@@" min="0">
        </label>
      </div>
      <div>
        <label>Ліміт постів на ключове слово
          <input type="number" name="limit" id="limitField" value="@@LIMIT@@" min="1">
          <button type="button" class="testbtn" onclick="document.getElementById('limitField').value=10">Тест: 10</button>
        </label>
      </div>
    </div>
  </fieldset>

  <fieldset id="scrapeOnly" style="display:none;">
    <legend>Параметри збору (тільки для повного режиму)</legend>
    <div class="row">
      <div><label>Глибина коментарів <input type="number" name="max_depth" value="@@MAX_DEPTH@@" min="0"></label></div>
      <div><label>Мін. рейтинг коментаря <input type="number" name="min_score" value="@@MIN_SCORE@@"></label></div>
    </div>
    <div class="row">
      <div><label>Затримка між постами (сек) <input type="number" name="delay" value="@@DELAY@@" min="1"></label></div>
      <div><label>Постів на 1 MD-файл <input type="number" name="posts_per_file" value="@@POSTS_PER_FILE@@" min="1"></label></div>
    </div>
  </fieldset>

  <fieldset>
    <legend>Доступ</legend>
    <label>Cookie сесії Reddit (reddit_session)
      <input type="password" name="cookie" required>
    </label>
    <div class="hint">Не зберігається на диску. Живе тільки поки застосунок працює.</div>
  </fieldset>

  <button type="submit">▶ Почати</button>
</form>

<p class="hint" style="margin-top:12px; margin-bottom:20px;">Налаштування (крім cookie) автоматично збережуться після старту й підставляться наступного разу.</p>

<script>
function toggleMode() {
  const mode = document.querySelector('input[name=mode]:checked').value;
  document.getElementById('scrapeOnly').style.display = (mode === 'scrape') ? 'block' : 'none';
}
document.getElementById('jobForm').addEventListener('submit', async (e) => {
  e.preventDefault();
  const fd = new FormData(e.target);
  const resp = await fetch('/start', { method: 'POST', body: fd });
  if (resp.ok) { window.location.href = '/progress'; }
  else { alert('Помилка запуску: ' + await resp.text()); }
});
</script>
</body>
</html>
"""

PROGRESS_HTML = """
<!DOCTYPE html>
<html lang="uk">
<head>
<meta charset="UTF-8">
<title>Виконання — Reddit Research Tool</title>
<style>
  :root {
    --bg: #101828;
    --card: #1A2438;
    --border: #2E3B57;
    --accent: #B8F500;
    --text: #FFFFFF;
    --text-secondary: #7A8499;
  }
  body { font-family: 'Calibri', 'Segoe UI', sans-serif; max-width: 800px; margin: 30px auto;
    padding: 0 20px; color: var(--text); background: var(--bg); }
  h1 { font-size: 20px; font-weight: 500; }
  #bar { background: var(--card); border-radius: 8px; height: 14px; overflow: hidden; margin: 14px 0; }
  #fill { background: var(--accent); height: 100%; width: 0%; transition: width 0.3s; }
  #log { background: var(--card); color: var(--text); font-family: 'Consolas', 'Calibri', monospace;
    font-size: 12px; padding: 12px 14px; height: 340px; overflow-y: auto; border-radius: 8px;
    white-space: pre-wrap; line-height: 1.6; }
  button { padding: 9px 18px; border-radius: 6px; border: none; font-size: 13px;
    font-family: 'Calibri', sans-serif; cursor: pointer; margin-right: 8px; }
  #pauseBtn { background: transparent; color: var(--accent); border: 1px solid var(--accent); }
  #stopBtn { background: transparent; color: var(--text); border: 1px solid var(--border); }
  #backBtn { background: transparent; color: var(--text-secondary); border: 1px solid var(--border); }
  #status { font-weight: 500; color: var(--accent); }
  #counter, #eta { color: var(--text-secondary); }
  #downloads a { display: inline-block; margin-top: 10px; padding: 10px 20px; background: var(--accent);
    color: var(--bg); border-radius: 8px; text-decoration: none; font-weight: 500; }
  .cookie-box { margin: 14px 0; padding: 14px 16px; background: var(--card); border-radius: 10px;
    border: 1px solid var(--border); }
  .cookie-box b { color: var(--text); font-size: 13px; }
  .cookie-box input { flex: 1; padding: 7px; background: var(--bg); color: var(--text);
    border: 1px solid var(--border); border-radius: 6px; font-family: 'Calibri', sans-serif; }
  .cookie-box button { background: var(--border); color: var(--accent); margin-top: 0; }
</style>
</head>
<body>
<h1>Виконання</h1>
<p>Статус: <span id="status">...</span> — <span id="counter"></span></p>
<div id="bar"><div id="fill"></div></div>

<button id="pauseBtn" onclick="pauseResume()">⏸ Пауза</button>
<button id="stopBtn" onclick="stopJob()">⏹ Стоп</button>
<button id="backBtn" onclick="location.href='/'">← Нові налаштування</button>

<p id="eta" style="font-size:13px;"></p>

<div class="cookie-box">
  <b>Оновити cookie сесії</b> <span style="color:var(--text-secondary); font-size:12px;">(якщо застосунок автоматично став на паузу через 403)</span>
  <div style="display:flex; gap:8px; margin-top:8px;">
    <input type="password" id="newCookie" placeholder="reddit_session">
    <button onclick="updateCookie()">Оновити й продовжити</button>
  </div>
</div>

<div id="downloads"></div>
<div id="log"></div>

<script>
let paused = false;
let startTime = Date.now();
let gotFirstProgress = false;

async function poll() {
  const r = await fetch('/status');
  const s = await r.json();

  document.getElementById('status').textContent = s.status;
  const unit = (s.mode === 'count') ? 'комбінацій' : 'постів';
  document.getElementById('counter').textContent = s.total
    ? (s.current + ' / ' + s.total + ' ' + unit + (s.comments ? ' • ' + s.comments.toLocaleString() + ' коментарів' : ''))
    : '';
  const pct = s.total ? Math.round(100 * s.current / s.total) : 0;
  document.getElementById('fill').style.width = pct + '%';

  if (s.current > 0 && s.total > 0) {
    if (!gotFirstProgress) { startTime = Date.now(); gotFirstProgress = true; }
    const elapsed = (Date.now() - startTime) / 1000;
    const rate = elapsed / s.current;
    const remaining = Math.round(rate * (s.total - s.current));
    const h = Math.floor(remaining / 3600);
    const m = Math.round((remaining % 3600) / 60);
    document.getElementById('eta').textContent = s.status === 'running'
      ? `Орієнтовно залишилось: ${h > 0 ? h + ' год ' : ''}${m} хв`
      : '';
  }

  const logEl = document.getElementById('log');
  logEl.textContent = s.log.join('\\n');
  logEl.scrollTop = logEl.scrollHeight;

  paused = (s.status === 'paused');
  document.getElementById('pauseBtn').textContent = paused ? '▶ Продовжити' : '⏸ Пауза';

  const dl = document.getElementById('downloads');
  dl.innerHTML = '';
  if (s.result_files && s.result_files.length) {
    s.result_files.forEach(f => {
      const a = document.createElement('a');
      a.href = '/download/' + encodeURIComponent(f);
      a.textContent = '⬇ Завантажити ' + f;
      dl.appendChild(a);
    });
  }

  if (s.status !== 'done' && s.status !== 'error') {
    setTimeout(poll, 1000);
  }
}

function pauseResume() {
  fetch(paused ? '/resume' : '/pause', { method: 'POST' }).then(poll);
}
function stopJob() {
  if (confirm('Зупинити виконання?')) {
    fetch('/stop', { method: 'POST' }).then(poll);
  }
}
function updateCookie() {
  const val = document.getElementById('newCookie').value.trim();
  if (!val) { alert('Встав значення cookie'); return; }
  const fd = new FormData();
  fd.append('cookie', val);
  fetch('/update_cookie', { method: 'POST', body: fd }).then(r => {
    if (r.ok) { document.getElementById('newCookie').value = ''; poll(); }
    else { alert('Помилка оновлення cookie'); }
  });
}

poll();
</script>
</body>
</html>
"""


def render_index() -> str:
    s = load_settings()

    def opt(select_options, current_value):
        parts = []
        for value, label in select_options:
            sel = " selected" if value == current_value else ""
            parts.append(f'<option value="{value}"{sel}>{label}</option>')
        return "\n            ".join(parts)

    sort_options = opt(
        [("new", "new (найновіші)"), ("relevance", "relevance"),
         ("top", "top"), ("comments", "comments")],
        s["sort"],
    )
    time_options = opt(
        [("all", "all"), ("year", "year"), ("month", "month"),
         ("week", "week"), ("day", "day")],
        s["time_filter"],
    )

    html_out = INDEX_HTML
    html_out = html_out.replace("@@BASE_NAME@@", html.escape(str(s["base_name"])))
    html_out = html_out.replace("@@MODE_COUNT_CHECKED@@", "checked" if s["mode"] == "count" else "")
    html_out = html_out.replace("@@MODE_SCRAPE_CHECKED@@", "checked" if s["mode"] == "scrape" else "")
    html_out = html_out.replace("@@KEYWORDS@@", html.escape(str(s["keywords"])))
    html_out = html_out.replace("@@SUBREDDITS@@", html.escape(str(s["subreddits"])))
    html_out = html_out.replace("@@SORT_OPTIONS@@", sort_options)
    html_out = html_out.replace("@@TIME_OPTIONS@@", time_options)
    html_out = html_out.replace("@@MONTHS_BACK@@", str(s["months_back"]))
    html_out = html_out.replace("@@LIMIT@@", str(s["limit"]))
    html_out = html_out.replace("@@MAX_DEPTH@@", str(s["max_depth"]))
    html_out = html_out.replace("@@MIN_SCORE@@", str(s["min_score"]))
    html_out = html_out.replace("@@DELAY@@", str(s["delay"]))
    html_out = html_out.replace("@@POSTS_PER_FILE@@", str(s["posts_per_file"]))
    return html_out


@app.route("/")
def index():
    return render_index()


@app.route("/start", methods=["POST"])
def start():
    with job_lock:
        if job["status"] == "running":
            return "Задача вже виконується", 400

    f = request.form
    keywords = [k.strip() for k in f.get("keywords", "").splitlines() if k.strip()]
    subreddits = [s.strip() for s in f.get("subreddits", "").splitlines() if s.strip()]

    if not keywords:
        return "Вкажи хоча б одне ключове слово", 400

    base_name = re.sub(r'[^\w\-]', '_', f.get("base_name", "reddit_project").strip()) or "reddit_project"

    params = {
        "mode": f.get("mode", "count"),
        "base_name": base_name,
        "keywords": keywords,
        "subreddits": subreddits,
        "sort": f.get("sort", "new"),
        "time_filter": f.get("time_filter", "all"),
        "months_back": int(f.get("months_back", 0) or 0),
        "limit": int(f.get("limit", 1000) or 1000),
        "cookie": f.get("cookie", "").strip(),
        "max_depth": int(f.get("max_depth", 6) or 6),
        "min_score": int(f.get("min_score", -10) or -10),
        "comment_limit": 500,
        "delay": int(f.get("delay", 5) or 5),
        "max_retries": 3,
        "posts_per_file": int(f.get("posts_per_file", 120) or 120),
    }

    if not params["cookie"]:
        return "Вкажи cookie сесії", 400
    if params["mode"] == "count" and not subreddits:
        return "Для режиму підрахунку потрібен хоча б один сабредіт", 400

    save_settings(params)

    pause_event.clear()
    stop_event.clear()

    t = threading.Thread(target=run_job_thread, args=(params,), daemon=True)
    t.start()
    return jsonify({"ok": True})


@app.route("/progress")
def progress_page():
    return PROGRESS_HTML


@app.route("/status")
def status():
    with job_lock:
        return jsonify({
            "status": job["status"],
            "current": job["current"],
            "total": job["total"],
            "comments": job["comments"],
            "mode": job["mode"],
            "log": job["log"][-400:],
            "result_files": job["result_files"],
        })


@app.route("/update_cookie", methods=["POST"])
def update_cookie():
    new_cookie = request.form.get("cookie", "").strip()
    if not new_cookie:
        return "Порожня cookie", 400
    session_cookie["reddit_session"] = new_cookie
    consecutive_403["count"] = 0
    pause_event.clear()
    log("🔄 Cookie оновлено. Продовжую роботу...")
    with job_lock:
        if job["status"] == "paused":
            job["status"] = "running"
    return jsonify({"ok": True})


@app.route("/pause", methods=["POST"])
def pause():
    pause_event.set()
    return jsonify({"ok": True})


@app.route("/resume", methods=["POST"])
def resume():
    pause_event.clear()
    return jsonify({"ok": True})


@app.route("/stop", methods=["POST"])
def stop():
    stop_event.set()
    pause_event.clear()  # щоб не застрягло на паузі
    return jsonify({"ok": True})


@app.route("/download/<path:filename>")
def download(filename):
    path = OUTPUT_DIR / filename
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True)


if __name__ == "__main__":
    print("=" * 50)
    print("Reddit Research Tool запущено.")
    print("Відкрий у браузері: http://localhost:5000")
    print("=" * 50)
    app.run(host="127.0.0.1", port=5000, debug=False)
